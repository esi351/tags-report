from openpyxl import load_workbook
from openpyxl.styles import Font
import os
import shutil
import jdatetime


# 📦 تابع بکاپ فقط از دو فایل خارجی مشخص‌شده
def backup_external_files():
    # محاسبه سال جاری به صورت داینامیک
    current_year = jdatetime.datetime.now().year
    now = jdatetime.datetime.now()
    
    # مسیر پوشه بکاپ (در کنار فایل ترانسمیتر)
    backup_root = r"Z:\electrical\8-Maintenance History\Backup"
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H-%M")
    dest_folder = os.path.join(backup_root, date_str, time_str)
    
    os.makedirs(dest_folder, exist_ok=True)

    # لیست فایل‌های خارجی با مسیر کامل
    files_to_backup = [
        rf"z:\ELECTRICAL\4-ELECTRICITY CONSUMPTION\ENERGY Report -24 o'clock-{current_year}.xlsb",
        r"z:\ELECTRICAL\5-Daily Report\Daily Report.xlsx"
    ]

    for src_path in files_to_backup:
        if os.path.exists(src_path):
            filename = os.path.basename(src_path)
            dst = os.path.join(dest_folder, filename)
            try:
                shutil.copy2(src_path, dst)
                print(f"[OK] بکاپ گرفته شد: {filename}")
            except Exception as e:
                print(f"[ERROR] مشکل در کپی {filename}: {str(e)}")
        else:
            print(f"[SKIP] فایل وجود ندارد: {src_path}")


#  تابع پاک کردن سلول‌ها
def clear_columns(sheet):
    for row in sheet.iter_rows(max_row=800, max_col=11):
        for cell in row:
            cell.value = None


# 📊 تابع پردازش فایل ALL TRANSMITER.xlsx
def process_transmitter_excel(file_path):
    # 🔁 مرحله ۱: بکاپ از دو فایل خارجی
    backup_external_files()

    # 🔁 مرحله ۲: بارگذاری فایل
    try:
        wb = load_workbook(file_path)
    except Exception as e:
        print(f"[ERROR] خطا در باز کردن فایل: {str(e)}")
        return

    if 'List' not in wb.sheetnames:
        print("Sheet 'List' not found!")
        return

    list_sheet = wb['List']
    clear_columns(list_sheet)

    # محاسبه سال جاری
    current_year = jdatetime.datetime.now().year
    emsal_str = str(current_year)

    # تعریف سربرگ‌ها
    headers = [
        "List", 
        "تاریخ آخرین کالیبره", 
        "تاریخ آخرین روتین B", 
        "اسامی انجام‌دهنده آخرین روتین",  
        "ساعت کاری آخرین روتین",           
        "تاریخ اعلام نیاز به خرید قطعه",
        "تاریخ آخرین کار تعمیراتی غیر روتین و کالیبره", 
        f"تعداد کالیبره سال {emsal_str}",
        f"تعداد روتین سال {emsal_str}", 
        f"تعداد کارهای تعمیراتی غیر روتین و کالیبره {emsal_str}",
        f"تعداد روتین 6 ماهه {emsal_str}"
    ]

    for col_idx, header in enumerate(headers, start=1):
        list_sheet.cell(row=1, column=col_idx).value = header

    row_idx = 2

    for sheet_name in wb.sheetnames:
        if sheet_name == 'List':
            continue

        sheet = wb[sheet_name]

        calibration_dates = []
        routine_dates = []
        routine_persons = []
        routine_hours = []
        repair_dates = []
        buy_flag = False

        for row in sheet.iter_rows(values_only=True):
            b_value = str(row[1]) if len(row) > 1 else ""
            e_value = row[4] if len(row) > 4 else None
            f_value = row[5] if len(row) > 5 else None
            g_value = row[6] if len(row) > 6 else None
            h_value = row[7] if len(row) > 7 else None
            i_value = row[8] if len(row) > 8 else None
            j_value = row[9] if len(row) > 9 else None

            if f_value == chr(252) and b_value.startswith("14"):
                calibration_dates.append(b_value)

            if e_value == chr(252) and b_value.startswith("14"):
                routine_dates.append(b_value)
                routine_persons.append(str(i_value) if i_value else "")
                try:
                    hour_val = float(j_value) if j_value is not None else 0
                except (ValueError, TypeError):
                    hour_val = 0
                routine_hours.append(hour_val)

            if g_value == chr(252):
                buy_flag = True

            if h_value == chr(252) and b_value.startswith("14"):
                repair_dates.append(b_value)

        cell = list_sheet.cell(row=row_idx, column=1)
        cell.value = sheet_name
        cell.hyperlink = f"#'{sheet_name}'!A1"
        cell.font = Font(color="0000FF", underline="single")

        if calibration_dates:
            list_sheet.cell(row=row_idx, column=2).value = max(calibration_dates)
        else:
            list_sheet.cell(row=row_idx, column=2).value = "-"

        if routine_dates:
            last_routine_date = max(routine_dates)
            list_sheet.cell(row=row_idx, column=3).value = last_routine_date

            matching_indices = [i for i, d in enumerate(routine_dates) if d == last_routine_date]
            persons_list = [routine_persons[i] for i in matching_indices if routine_persons[i]]
            persons_str = "::::".join(persons_list) if persons_list else "-"
            total_hours = sum(routine_hours[i] for i in matching_indices)

            list_sheet.cell(row=row_idx, column=4).value = persons_str
            list_sheet.cell(row=row_idx, column=5).value = total_hours
        else:
            list_sheet.cell(row=row_idx, column=3).value = "-"
            list_sheet.cell(row=row_idx, column=4).value = "-"
            list_sheet.cell(row=row_idx, column=5).value = 0

        list_sheet.cell(row=row_idx, column=6).value = "Yes" if buy_flag else "-"

        if repair_dates:
            list_sheet.cell(row=row_idx, column=7).value = max(repair_dates)
        else:
            list_sheet.cell(row=row_idx, column=7).value = "-"

        calibration_current_year = sum(1 for x in calibration_dates if x.startswith(emsal_str))
        routine_current_year = sum(1 for x in routine_dates if x.startswith(emsal_str))
        repair_current_year = sum(1 for x in repair_dates if x.startswith(emsal_str))

        current_month = jdatetime.datetime.now().month
        if 1 <= current_month <= 6:
            routine_last_6_months = sum(1 for x in routine_dates if x.startswith(emsal_str) and 1 <= int(x[5:7]) <= 6)
        elif 7 <= current_month <= 12:
            routine_last_6_months = sum(1 for x in routine_dates if x.startswith(emsal_str) and 7 <= int(x[5:7]) <= 12)
        else:
            routine_last_6_months = 0

        list_sheet.cell(row=row_idx, column=8).value = calibration_current_year
        list_sheet.cell(row=row_idx, column=9).value = routine_current_year
        list_sheet.cell(row=row_idx, column=10).value = repair_current_year
        list_sheet.cell(row=row_idx, column=11).value = routine_last_6_months

        row_idx += 1

    # 🔚 تمیزکاری نهایی
    max_row = list_sheet.max_row
    max_col = list_sheet.max_column

    if max_row >= 405:
        list_sheet.delete_rows(405, max_row - 404)
    if max_col >= 110:
        list_sheet.delete_cols(110, max_col - 109)

    wb.save(file_path)
    print(f"✅ پردازش تکمیل و فایل ذخیره شد: {file_path}")


# 📍 مسیر فایل اصلی
file_path = r"Z:\electrical\8-Maintenance History\ALL TRANSMITER.xlsx"

if __name__ == "__main__":
    process_transmitter_excel(file_path)
